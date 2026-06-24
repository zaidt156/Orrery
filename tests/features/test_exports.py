import io
import zipfile

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.features import exports

SAMPLE_REPLY = """# Quarterly report

Revenue increased by **12%**.

- North region improved
- South region stayed flat

```python
print("trusted text only")
```

| Metric | Value |
|---|---:|
| Formula-like text | =2+2 |
| Revenue | 120 |
"""


def test_parse_markdown_preserves_document_blocks():
    blocks = exports.parse_markdown(SAMPLE_REPLY)
    kinds = [block.kind for block in blocks]

    assert "heading" in kinds
    assert "bullet" in kinds
    assert "code" in kinds
    assert "table" in kinds


def test_pdf_export_is_readable():
    result = exports.render_export("Quarterly report", "openai/test", SAMPLE_REPLY, "pdf")

    assert result.content.startswith(b"%PDF")
    text = "\n".join(
        page.extract_text() or "" for page in PdfReader(io.BytesIO(result.content)).pages
    )
    assert "Quarterly report" in text
    assert "Revenue increased" in text


def test_docx_export_is_readable():
    result = exports.render_export("Quarterly report", "openai/test", SAMPLE_REPLY, "docx")
    document = Document(io.BytesIO(result.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert result.content.startswith(b"PK")
    assert "Quarterly report" in text
    assert "trusted text only" in text


def test_xlsx_export_neutralizes_formula_cells():
    result = exports.render_export("Quarterly report", "openai/test", SAMPLE_REPLY, "xlsx")
    workbook = load_workbook(io.BytesIO(result.content), data_only=False)

    assert result.content.startswith(b"PK")
    assert workbook["Reply"]["B2"].value == "Quarterly report"
    assert workbook["Table 1"]["B2"].value == "'=2+2"
    assert exports._excel_safe("  @SUM(A1:A2)") == "'  @SUM(A1:A2)"


def test_export_rejects_oversized_reply():
    with pytest.raises(exports.ExportTooLarge):
        exports.parse_markdown("x" * (exports.MAX_EXPORT_CHARS + 1))



def test_requested_formats_detects_only_asked_files():
    assert exports.requested_formats("Create a PowerPoint and CSV for this data") == ["pptx", "csv"]
    assert exports.requested_formats("Explain this normally") == []


def test_select_export_content_removes_chat_wrapper_for_file():
    reply = """Here is the CSV file:

| Name | Value |
|---|---:|
| A | 1 |

Let me know if you want changes."""
    cleaned = exports.select_export_content("make a csv", reply, "csv")

    assert cleaned.startswith("| Name | Value |")
    assert "Here is" not in cleaned
    assert "Let me know" not in cleaned


def test_csv_export_uses_requested_table_only():
    content = """Here is the spreadsheet:

| Item | Amount |
|---|---:|
| Total | =2+2 |

Extra explanation that should not be part of the sheet."""
    cleaned = exports.select_export_content("export as csv", content, "csv")
    result = exports.render_export("Data", "openai/test", cleaned, "csv")

    assert result.filename.endswith(".csv")
    body = result.content.decode("utf-8-sig")
    assert "Item,Amount" in body
    assert "'=2+2" in body
    assert "Extra explanation" not in body


def test_json_export_is_valid_even_with_fenced_json():
    cleaned = exports.select_export_content("make a json file", '```json\n{"ok": true}\n```', "json")
    result = exports.render_export("Payload", "openai/test", cleaned, "json")

    assert result.filename.endswith(".json")
    assert '"ok": true' in result.content.decode("utf-8")


def test_text_and_markdown_exports_are_plain_files():
    txt = exports.render_export("Notes", "openai/test", "# Title\n\nBody", "txt")
    md = exports.render_export("Notes", "openai/test", "# Title\n\nBody", "md")

    assert txt.media_type.startswith("text/plain")
    assert txt.content.decode("utf-8").startswith("Title")
    assert md.media_type.startswith("text/markdown")
    assert md.content.decode("utf-8").startswith("# Title")


def test_pptx_export_contains_office_parts():
    result = exports.render_export("Deck", "openai/test", "# Slide one\n\n- Point A\n- Point B", "pptx")

    assert result.content.startswith(b"PK")
    with zipfile.ZipFile(io.BytesIO(result.content)) as zf:
        names = set(zf.namelist())
        assert "ppt/presentation.xml" in names
        assert "ppt/slides/slide1.xml" in names
        slide = zf.read("ppt/slides/slide1.xml").decode("utf-8")
    assert "Slide one" in slide
    assert "Point A" in slide


def test_preview_for_non_pdf_is_html():
    result = exports.render_preview("Preview", "openai/test", "| A | B |\n|---|---|\n| 1 | 2 |", "xlsx")

    assert result.media_type.startswith("text/html")
    assert b"XLSX preview" in result.content
